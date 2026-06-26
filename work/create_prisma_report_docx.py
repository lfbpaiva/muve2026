from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "RELATORIO_PRISMA_ORM_VS_PRISMA_POSTGRES.md"
OUT_DIR = ROOT / ".docs"
OUT = OUT_DIR / "RELATORIO_PRISMA_ORM_VS_PRISMA_POSTGRES.docx"


COLORS = {
    "blue": RGBColor(0x2E, 0x74, 0xB5),
    "dark_blue": RGBColor(0x1F, 0x4D, 0x78),
    "ink": RGBColor(0x1F, 0x29, 0x37),
    "muted": RGBColor(0x55, 0x65, 0x72),
    "light_fill": "F2F4F7",
    "callout_fill": "F4F6F9",
    "border": "D9E2EC",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline_markdown(paragraph, text, *, size=11, color=None, bold_default=False):
    tokens = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", size - 0.5, RGBColor(0x2F, 0x3A, 0x45))
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, "Calibri", size, color, bold=True)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, "Calibri", size, color, bold=bold_default)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = COLORS["ink"]
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, COLORS["blue"], 16, 8),
        ("Heading 2", 13, COLORS["blue"], 12, 6),
        ("Heading 3", 12, COLORS["dark_blue"], 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_footer(doc):
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("MUVE | Relatorio de decisao Prisma")
        set_run_font(run, "Calibri", 9, COLORS["muted"])


def add_title_page(doc):
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("Relatorio de decisao: Prisma ORM vs Prisma Postgres")
    set_run_font(title_run, "Calibri", 22, COLORS["blue"], bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    add_inline_markdown(
        subtitle,
        "Projeto MUVE | Analise tecnica para decisao da equipe | 2026-06-09",
        size=11,
        color=COLORS["muted"],
    )

    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    set_table_width(callout, [9360])
    set_cell_margins(callout, top=120, bottom=120, start=160, end=160)
    cell = callout.cell(0, 0)
    set_cell_shading(cell, COLORS["callout_fill"])
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_markdown(
        p,
        "Recomendacao: manter Prisma ORM no backend. Tratar Prisma Postgres como uma decisao posterior de hospedagem do banco, nao como substituto do ORM.",
        size=11,
        color=COLORS["ink"],
        bold_default=True,
    )
    doc.add_paragraph()


def table_widths_for(rows):
    cols = len(rows[0])
    if cols == 2:
        return [2500, 6860]
    if cols == 3:
        return [1800, 3780, 3780]
    if cols == 4:
        return [2100, 1450, 2905, 2905]
    return [int(9360 / cols)] * cols


def add_markdown_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    widths = table_widths_for(rows)
    set_table_width(table, widths)
    set_cell_margins(table)

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if r_idx == 0:
                set_cell_shading(cell, COLORS["light_fill"])
                add_inline_markdown(p, value, size=10.5, color=COLORS["ink"], bold_default=True)
            else:
                add_inline_markdown(p, value, size=10, color=COLORS["ink"])
    doc.add_paragraph()


def parse_table(lines, start):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        raw = lines[idx].strip()
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def add_code_block(doc, code_lines, language):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [9360])
    set_cell_margins(table, top=100, bottom=100, start=160, end=160)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F9FB")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if language:
        lang_run = p.add_run(f"{language}\n")
        set_run_font(lang_run, "Calibri", 9, COLORS["muted"], bold=True)
    for line_idx, line in enumerate(code_lines):
        run = p.add_run(line)
        set_run_font(run, "Consolas", 9, RGBColor(0x2F, 0x3A, 0x45))
        if line_idx != len(code_lines) - 1:
            p.add_run("\n")
    doc.add_paragraph()


def add_list_item(doc, text, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    add_inline_markdown(p, text, size=11)


def build_doc():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_title_page(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# "):
            # Already rendered a stronger title page.
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            i += 1
            continue

        if stripped.startswith("```"):
            language = stripped.strip("`").strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines, language)
            i += 1
            continue

        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                add_markdown_table(doc, rows)
            continue

        if stripped.startswith(">"):
            table = doc.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            set_table_width(table, [9360])
            set_cell_margins(table, top=120, bottom=120, start=160, end=160)
            cell = table.cell(0, 0)
            set_cell_shading(cell, COLORS["callout_fill"])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_markdown(p, stripped.lstrip("> ").strip(), size=11, bold_default=True)
            doc.add_paragraph()
            i += 1
            continue

        numbered_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered_match:
            add_list_item(doc, numbered_match.group(1), numbered=True)
            i += 1
            continue

        if stripped.startswith("- "):
            add_list_item(doc, stripped[2:], numbered=False)
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_markdown(p, stripped, size=11)
        i += 1

    add_footer(doc)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_doc())
